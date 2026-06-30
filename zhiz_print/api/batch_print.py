# -*- coding: utf-8 -*-
# Copyright (c) 2026, Zhiz Print
# Batch printing API endpoints

from __future__ import unicode_literals
import frappe
from frappe import _
import json


def _auto_match_designs_for_docs(doctype, docnames):
    """Match each docname to the first design whose enable_condition passes.
    Returns dict: {docname: design_info_dict} and set of all matched design names."""
    from zhiz_print.zhiz_print.doctype.super_print_design.super_print_design import SuperPrintDesign

    designs = frappe.get_all(
        "Super Print Design",
        filters={"target_doctype": doctype, "enabled": 1},
        fields=["name", "design_name", "print_paper", "priority", "orientation"],
        order_by="priority asc, design_name",
    )

    if not designs:
        return {}, set()

    # Preload paper info for all designs
    design_map = {}
    for d in designs:
        if d.print_paper:
            paper = frappe.db.get_value(
                "Super Print Paper", d.print_paper,
                ["width", "height", "margin_top", "margin_bottom", "margin_left", "margin_right"],
                as_dict=True,
            )
            if paper:
                d["paper_width"] = paper.width
                d["paper_height"] = paper.height
                d["margin_top"] = paper.margin_top or 0
                d["margin_bottom"] = paper.margin_bottom or 0
                d["margin_left"] = paper.margin_left or 0
                d["margin_right"] = paper.margin_right or 0
        # v15.04.33: orientation for batch rendering
        d["orientation"] = d.get("orientation") or "Auto"
        design_map[d.name] = d

    doc_matches = {}
    matched_designs = set()

    for docname in docnames:
        try:
            doc = frappe.get_doc(doctype, docname)
        except Exception:
            continue

        for d in designs:
            if SuperPrintDesign.check_enable_conditions(d.name, doc=doc):
                doc_matches[docname] = d
                matched_designs.add(d.name)
                break

    return doc_matches, matched_designs


@frappe.whitelist()
def check_batch_print_enabled(doctype):
    """Check if super print is enabled for a given doctype and return available designs."""
    setting = frappe.get_single("Zprint Setting")
    if not frappe.utils.cint(setting.get("enable_super_print_page")):
        return {"enabled": False, "designs": []}

    enable_mode = setting.get("print_enable_mode") or "Enable for All"
    if enable_mode == "Enable for Specific":
        enabled_doctypes = [
            item.doctype_name for item in setting.get("print_enabled_doctypes", [])
            if item.enabled and item.doctype_name
        ]
        if doctype not in enabled_doctypes:
            return {"enabled": False, "designs": []}

    designs = frappe.get_all(
        "Super Print Design",
        filters={"target_doctype": doctype, "enabled": 1},
        fields=["name", "design_name", "print_paper", "priority"],
        order_by="priority asc, design_name",
    )

    for d in designs:
        if d.print_paper:
            paper = frappe.db.get_value(
                "Super Print Paper",
                d.print_paper,
                ["width", "height", "margin_top", "margin_bottom", "margin_left", "margin_right"],
                as_dict=True,
            )
            if paper:
                d["paper_width"] = paper.width
                d["paper_height"] = paper.height
                d["margin_top"] = paper.margin_top or 0
                d["margin_bottom"] = paper.margin_bottom or 0
                d["margin_left"] = paper.margin_left or 0
                d["margin_right"] = paper.margin_right or 0

        params_count = frappe.db.count("Super Print Design Parameter", {"parent": d.name})
        d["has_parameters"] = params_count > 0
        if d.has_parameters:
            d["parameters"] = frappe.get_all(
                "Super Print Design Parameter",
                filters={"parent": d.name},
                fields=["param_name", "param_label", "param_type", "default_value", "reqd", "options"],
            )

    return {"enabled": True, "designs": designs}


@frappe.whitelist()
def batch_render_preview(doctype, docnames, design_name=None, params=None, auto_match=False):
    """Render preview HTML for multiple documents."""
    from zhiz_print.api.print_designer import _check_license, _render_print_html, _is_draft_blocked

    _check_license()

    if isinstance(docnames, str):
        docnames = json.loads(docnames)
    if isinstance(params, str):
        try:
            params = json.loads(params)
        except (json.JSONDecodeError, TypeError):
            params = {}
    auto_match = frappe.utils.cint(auto_match)

    results = []
    errors = []
    skipped = []

    if auto_match:
        doc_matches, _matched_set = _auto_match_designs_for_docs(doctype, docnames)

        for docname in docnames:
            matched = doc_matches.get(docname)
            if not matched:
                errors.append({"docname": docname, "error": "No matching template found"})
                continue
            dname = matched["name"]
            # v15.04.27: skip drafts when design has draft_no_print=1 — preview, PDF, Excel all
            # apply the same per-doc policy so what you see matches what you can export.
            try:
                if _is_draft_blocked(dname, frappe.get_doc(doctype, docname)):
                    skipped.append({
                        "docname": docname,
                        "design_name": dname,
                        "design_label": matched.get("design_name", dname),
                        "reason": "draft",
                    })
                    continue
            except Exception as e:
                frappe.log_error(f"Batch draft check failed for {doctype} {docname}: {e}")
                errors.append({"docname": docname, "error": str(e)})
                continue
            try:
                # v15.04.33: preview uses original paper W×H — pass 'Auto' so the
                # design's saved orientation is NOT applied to preview HTML.
                # Browser print path reads `orientation` field and rotates via CSS.
                html, _ctx = _render_print_html(
                    doctype, docname, dname, params,
                    skip_px_scaling=True, orientation_override="Auto",
                )
                results.append({
                    "docname": docname,
                    "html": html,
                    "design_name": dname,
                    "design_label": matched.get("design_name", dname),
                    "paper_width": matched.get("paper_width", 210),
                    "paper_height": matched.get("paper_height", 297),
                    "margin_top": matched.get("margin_top", 0),
                    "margin_bottom": matched.get("margin_bottom", 0),
                    "margin_left": matched.get("margin_left", 0),
                    "margin_right": matched.get("margin_right", 0),
                    "orientation": matched.get("orientation", "Auto"),
                })
            except Exception as e:
                frappe.log_error(f"Batch render failed for {doctype} {docname}: {e}")
                errors.append({"docname": docname, "error": str(e)})
    else:
        if not design_name:
            frappe.throw(_("Design name is required"))
        design = frappe.get_doc("Super Print Design", design_name)
        paper_width, paper_height = 210, 297
        margin_top = margin_bottom = margin_left = margin_right = 0

        if design.print_paper:
            paper = frappe.get_doc("Super Print Paper", design.print_paper)
            paper_width = paper.width
            paper_height = paper.height
            margin_top = paper.margin_top or 0
            margin_bottom = paper.margin_bottom or 0
            margin_left = paper.margin_left or 0
            margin_right = paper.margin_right or 0

        for docname in docnames:
            # v15.04.27: skip drafts per-doc — submitted docs proceed, drafts are listed in `skipped`.
            try:
                if _is_draft_blocked(design_name, frappe.get_doc(doctype, docname)):
                    skipped.append({
                        "docname": docname,
                        "design_name": design_name,
                        "design_label": design.design_name,
                        "reason": "draft",
                    })
                    continue
            except Exception as e:
                frappe.log_error(f"Batch draft check failed for {doctype} {docname}: {e}")
                errors.append({"docname": docname, "error": str(e)})
                continue
            try:
                # v15.04.33: preview uses original paper W×H — pass 'Auto' so the
                # design's saved orientation is NOT applied to preview HTML.
                html, _ctx = _render_print_html(
                    doctype, docname, design_name, params,
                    skip_px_scaling=True, orientation_override="Auto",
                )
                results.append({
                    "docname": docname,
                    "html": html,
                    "design_name": design_name,
                    "design_label": design.design_name,
                    "paper_width": paper_width,
                    "paper_height": paper_height,
                    "margin_top": margin_top,
                    "margin_bottom": margin_bottom,
                    "margin_left": margin_left,
                    "margin_right": margin_right,
                    "orientation": design.orientation or "Auto",
                })
            except Exception as e:
                frappe.log_error(f"Batch render failed for {doctype} {docname}: {e}")
                errors.append({"docname": docname, "error": str(e)})

    return {"results": results, "errors": errors, "skipped": skipped}


@frappe.whitelist()
def batch_generate_pdf(doctype, docnames, design_name=None, params=None, auto_match=False):
    """Generate merged PDF for multiple documents by concatenating HTML first."""
    from zhiz_print.api.print_designer import (
        _check_license, _check_draft_no_print, _is_draft_blocked,
        _render_print_html, _pdf_response,
        _fix_merged_cell_borders_for_pdf, _prepare_html_for_wkhtmltopdf,
    )

    _check_license()

    if isinstance(docnames, str):
        docnames = json.loads(docnames)
    if isinstance(params, str):
        try:
            params = json.loads(params)
        except (json.JSONDecodeError, TypeError):
            params = {}
    auto_match = frappe.utils.cint(auto_match)

    engine_mode = frappe.db.get_single_value("Zprint Setting", "pdf_engine_mode") or "wkhtmltopdf"
    success_count = 0
    html_parts = []
    styles_collected = set()

    import re

    if auto_match:
        doc_matches, _matched_set = _auto_match_designs_for_docs(doctype, docnames)
        # v15.04.25: draft_no_print now blocks per-doc (docstatus=0), not per-design.
        # Filter doc_matches so that any (doc, design) pair where the doc is draft
        # AND the design has draft_no_print=1 is dropped. Submitted docs keep all designs.
        doc_matches = {
            dn: info for dn, info in doc_matches.items()
            if not _is_draft_blocked(info["name"], frappe.get_doc(doctype, dn))
        }
        if not doc_matches:
            frappe.throw(_("All matched documents are in draft state and cannot be printed. Submit them first."), frappe.PermissionError)

        for idx, docname in enumerate(docnames):
            matched = doc_matches.get(docname)
            if not matched:
                continue
            try:
                dname = matched["name"]
                # v15.04.33: batch PDF does NOT rotate — force 'Auto' so the
                # design's saved orientation is NOT applied to PDF HTML.
                html, _ctx = _render_print_html(
                    doctype, docname, dname, params,
                    orientation_override="Auto",
                )
                body_match = re.search(r'<body[^>]*>([\s\S]*)</body>', html, re.IGNORECASE)
                style_matches = re.findall(r'<style[^>]*>[\s\S]*?</style>', html, re.IGNORECASE)

                body_content = body_match.group(1) if body_match else html
                if idx < len(docnames) - 1:
                    body_content += '<div style="page-break-after:always"></div>'

                for s in style_matches:
                    if s not in styles_collected:
                        styles_collected.add(s)

                html_parts.append(body_content)
                success_count += 1
            except Exception as e:
                frappe.log_error(f"Batch PDF render failed for {doctype} {docname}: {e}")

        if success_count == 0:
            frappe.throw(_("All documents failed to generate PDF"))

        combined_html = '<!DOCTYPE html>\n<html>\n<head>\n<meta charset="utf-8">\n'
        combined_html += "\n".join(styles_collected)
        combined_html += '\n</head>\n<body>\n'
        combined_html += "\n".join(html_parts)
        combined_html += '\n</body>\n</html>'

        # Use first matched design's paper for PDF settings
        first_matched = next(iter(doc_matches.values()), None)
        design = frappe.get_doc("Super Print Design", first_matched["name"]) if first_matched else None
    else:
        if not design_name:
            frappe.throw(_("Design name is required"))
        # v15.04.25: per-doc draft check moved into loop — submitted docs proceed,
        # draft docs are skipped individually so users can mix statuses in a batch.
        design = frappe.get_doc("Super Print Design", design_name)

        for idx, docname in enumerate(docnames):
            try:
                if _is_draft_blocked(design_name, frappe.get_doc(doctype, docname)):
                    continue
                # v15.04.33: batch PDF does NOT rotate — force 'Auto'.
                html, _ctx = _render_print_html(
                    doctype, docname, design_name, params,
                    orientation_override="Auto",
                )
                body_match = re.search(r'<body[^>]*>([\s\S]*)</body>', html, re.IGNORECASE)
                style_matches = re.findall(r'<style[^>]*>[\s\S]*?</style>', html, re.IGNORECASE)

                body_content = body_match.group(1) if body_match else html
                if idx < len(docnames) - 1:
                    body_content += '<div style="page-break-after:always"></div>'

                for s in style_matches:
                    if s not in styles_collected:
                        styles_collected.add(s)

                html_parts.append(body_content)
                success_count += 1
            except Exception as e:
                frappe.log_error(f"Batch PDF render failed for {doctype} {docname}: {e}")

        if success_count == 0:
            frappe.throw(_("All documents are in draft state and cannot be printed. Submit them first."))

        combined_html = '<!DOCTYPE html>\n<html>\n<head>\n<meta charset="utf-8">\n'
        combined_html += "\n".join(styles_collected)
        combined_html += '\n</head>\n<body>\n'
        combined_html += "\n".join(html_parts)
        combined_html += '\n</body>\n</html>'

    # v15.04.33: batch PDF does NOT apply orientation rotation — only browser
    # print path rotates. PDF engines receive the design's original paper size.

    # Generate single PDF using the selected engine
    if engine_mode == "WeasyPrint":
        from weasyprint import HTML as WeasyHTML
        combined_html = _fix_merged_cell_borders_for_pdf(combined_html)
        pdf_bytes = WeasyHTML(string=combined_html).write_pdf()
    elif engine_mode == "Chromium":
        pdf_bytes = _generate_chromium_pdf(combined_html, design)
    else:
        import pdfkit
        combined_html = _prepare_html_for_wkhtmltopdf(combined_html)
        options = {
            "quiet": "", "encoding": "UTF-8", "print-media-type": "",
            "background": "", "images": "", "disable-smart-shrinking": "",
            "margin-top": "0", "margin-bottom": "0",
            "margin-left": "0", "margin-right": "0",
        }
        if design and design.print_paper:
            paper = frappe.get_doc("Super Print Paper", design.print_paper)
            options["page-width"] = f"{paper.width}mm"
            options["page-height"] = f"{paper.height}mm"
        pdf_bytes = pdfkit.from_string(combined_html, False, options=options)

    label = "auto-match" if auto_match else (design.design_name if design else "batch")
    _pdf_response(pdf_bytes, f"batch-{doctype}-{success_count}docs-{label}.pdf")


def _generate_chromium_pdf(html, design):
    """Generate PDF via Chromium headless, return bytes."""
    import os
    import subprocess

    chromium_cmd = None
    for cmd in ['/snap/bin/chromium', 'chromium-browser', 'chromium', 'google-chrome', 'google-chrome-stable']:
        try:
            result = subprocess.run(['which', cmd], capture_output=True, text=True, timeout=5)
            if result.returncode == 0 and result.stdout.strip():
                chromium_cmd = result.stdout.strip()
                break
        except Exception:
            continue

    if not chromium_cmd:
        frappe.throw(_("Chromium is not installed"))

    tmpdir = os.path.abspath(os.path.join(frappe.get_site_path(), 'public', 'files', 'pdf_debug'))
    os.makedirs(tmpdir, exist_ok=True)
    html_path = os.path.join(tmpdir, 'chrome_batch_combined.html')
    pdf_path = os.path.join(tmpdir, 'chrome_batch_output.pdf')

    with open(html_path, 'w', encoding='utf-8') as f:
        f.write(html)

    if os.path.exists(pdf_path):
        os.remove(pdf_path)

    chrome_args = [
        chromium_cmd, '--headless=new', '--disable-gpu', '--no-sandbox',
        '--print-to-pdf=' + pdf_path, '--no-pdf-header-footer',
    ]
    if design.print_paper:
        paper = frappe.get_doc('Super Print Paper', design.print_paper)
        chrome_args.append('--print-to-pdf-options=' + json.dumps({
            'paperWidth': round(paper.width / 25.4, 4),
            'paperHeight': round(paper.height / 25.4, 4),
            'marginTop': 0, 'marginBottom': 0, 'marginLeft': 0, 'marginRight': 0,
        }))
    chrome_args.append('file://' + html_path)
    subprocess.run(chrome_args, capture_output=True, text=True, timeout=60)

    if not os.path.exists(pdf_path):
        frappe.throw(_("Chromium PDF generation failed"))

    with open(pdf_path, 'rb') as f:
        return f.read()


@frappe.whitelist()
def batch_export_excel(doctype, docnames, design_name=None, params=None, auto_match=False):
    """Export multi-sheet Excel for multiple documents."""
    from zhiz_print.api.print_designer import _check_license, _check_draft_no_print, _is_draft_blocked
    from io import BytesIO
    import openpyxl
    from zhiz_print.api.print_designer import _extract_css_rules, _write_table_to_excel, _set_column_widths

    _check_license()

    if isinstance(docnames, str):
        docnames = json.loads(docnames)
    if isinstance(params, str):
        try:
            params = json.loads(params)
        except (json.JSONDecodeError, TypeError):
            params = {}
    auto_match = frappe.utils.cint(auto_match)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    errors = []

    if auto_match:
        doc_matches, _matched_set = _auto_match_designs_for_docs(doctype, docnames)
        # v15.04.25: per-doc draft check (docstatus=0 + draft_no_print=1 → blocked).
        doc_matches = {
            dn: info for dn, info in doc_matches.items()
            if not _is_draft_blocked(info["name"], frappe.get_doc(doctype, dn))
        }

        for idx, docname in enumerate(docnames):
            matched = doc_matches.get(docname)
            if not matched:
                errors.append({"docname": docname, "error": "No matching template found"})
                continue
            try:
                if not frappe.has_permission(doctype, "print", docname):
                    errors.append({"docname": docname, "error": "No print permission"})
                    continue

                design = frappe.get_doc("Super Print Design", matched["name"])
                html = design.get_preview_for_document(doc_name=docname, params=params)
                _write_doc_to_sheet(wb, html, docname)
            except Exception as e:
                frappe.log_error(f"Batch Excel failed for {doctype} {docname}: {e}")
                errors.append({"docname": docname, "error": str(e)})

        label = "auto-match"
    else:
        if not design_name:
            frappe.throw(_("Design name is required"))
        # v15.04.25: per-doc draft check moved into loop.
        design = frappe.get_doc("Super Print Design", design_name)

        for idx, docname in enumerate(docnames):
            try:
                if not frappe.has_permission(doctype, "print", docname):
                    errors.append({"docname": docname, "error": "No print permission"})
                    continue

                if _is_draft_blocked(design_name, frappe.get_doc(doctype, docname)):
                    errors.append({"docname": docname, "error": "Document is in draft state"})
                    continue

                html = design.get_preview_for_document(doc_name=docname, params=params)
                _write_doc_to_sheet(wb, html, docname)
            except Exception as e:
                frappe.log_error(f"Batch Excel failed for {doctype} {docname}: {e}")
                errors.append({"docname": docname, "error": str(e)})

        label = design.design_name

    # v15.04.28: guard against empty workbook — happens when every doc was filtered
    # (drafts) or skipped (no permission / no template). Without this guard openpyxl
    # raises "At least one sheet must be visible" on save.
    if len(wb.worksheets) == 0:
        frappe.throw(_("No documents can be exported. All selected documents are in draft state or have no matching template."))

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    frappe.local.response.filename = f"batch-{doctype}-{len(docnames)}docs-{label}.xlsx"
    frappe.local.response.filecontent = output.getvalue()
    frappe.local.response.type = "binary"


def _write_doc_to_sheet(wb, html, docname, css_rules=None):
    """Render one document's HTML into a new openpyxl sheet.

    Extracted from batch_export_excel so the auto_match and named-design branches
    share the same sheet-building logic. css_rules is computed from the HTML when
    not supplied (callers in batch_export_excel pass None).
    """
    from bs4 import BeautifulSoup
    from zhiz_print.api.print_designer import _extract_css_rules, _write_table_to_excel, _set_column_widths

    soup = BeautifulSoup(html, 'html.parser')
    if css_rules is None:
        css_rules = _extract_css_rules(soup)

    sheet_name = str(docname)[:31]
    ws = wb.create_sheet(title=sheet_name)

    pages = []
    for page_div in soup.find_all('div', class_='print-page'):
        content_div = page_div.find('div', class_='print-page-content')
        if not content_div:
            continue
        table = content_div.find('table', class_='print-form-table')
        if table:
            pages.append(table)

    if not pages:
        ws.cell(row=1, column=1, value="No Data")
        return

    header_count = 0
    if len(pages) > 1:
        rows_1 = [tr.get_text(strip=True) for tr in pages[0].find_all('tr')]
        rows_2 = [tr.get_text(strip=True) for tr in pages[1].find_all('tr')]
        for i in range(min(len(rows_1), len(rows_2))):
            if rows_1[i] == rows_2[i]:
                header_count = i + 1
            else:
                break

    current_row = 1
    for page_idx, table in enumerate(pages):
        skip = header_count if page_idx > 0 else 0
        if page_idx == 0:
            _set_column_widths(ws, table, css_rules)
        rows_written = _write_table_to_excel(ws, table, current_row, css_rules, skip_rows=skip)
        current_row += rows_written


@frappe.whitelist()
def batch_record_print_log(doctype, docnames, design_name=None, params=None, export_type='Print', auto_match=False):
    """Record individual print logs for each document in the batch."""
    from zhiz_print.api.print_designer import _check_license

    _check_license()

    if isinstance(docnames, str):
        docnames = json.loads(docnames)
    if isinstance(params, str):
        try:
            params = json.loads(params)
        except (json.JSONDecodeError, TypeError):
            params = {}
    auto_match = frappe.utils.cint(auto_match)

    # Resolve per-doc design_name for auto_match
    doc_design_map = {}
    if auto_match:
        doc_matches, _matched_set = _auto_match_designs_for_docs(doctype, docnames)
        for docname, matched in doc_matches.items():
            doc_design_map[docname] = matched["name"]

    results = []
    for docname in docnames:
        try:
            dname = doc_design_map.get(docname, design_name)
            existing_count = frappe.db.count("Super Print Log", filters={
                "reference_doctype": doctype,
                "reference_name": docname,
            })
            log = frappe.get_doc({
                "doctype": "Super Print Log",
                "reference_doctype": doctype,
                "reference_name": docname,
                "print_design": dname,
                "export_type": export_type,
                "print_count": existing_count + 1,
                "print_user": frappe.session.user,
            })
            log.insert(ignore_permissions=True)
            results.append({
                "docname": docname,
                "log_name": log.name,
                "print_count": existing_count + 1,
            })
        except Exception as e:
            frappe.log_error(f"Batch log failed for {doctype} {docname}: {e}")

    frappe.db.commit()
    return {"results": results}


@frappe.whitelist()
def batch_get_print_logs(doctype, docnames):
    """Get aggregated print logs for multiple documents."""
    from zhiz_print.api.print_designer import _check_license

    _check_license()

    if isinstance(docnames, str):
        docnames = json.loads(docnames)

    if not docnames:
        return {"logs": [], "total_count": 0}

    filters = {"reference_doctype": doctype, "reference_name": ["in", docnames]}
    total_count = frappe.db.count("Super Print Log", filters=filters)
    logs = frappe.get_all(
        "Super Print Log",
        filters=filters,
        fields=["name", "reference_name", "print_design", "print_time", "print_user", "print_count", "export_type"],
        order_by="print_time desc",
        limit=100,
    )

    for log in logs:
        if log.print_user:
            log.user_fullname = frappe.db.get_value("User", log.print_user, "full_name") or log.print_user
        else:
            log.user_fullname = ''

    return {"logs": logs, "total_count": total_count}
